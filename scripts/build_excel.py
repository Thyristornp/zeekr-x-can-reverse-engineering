#!/usr/bin/env python3
"""Build the English workbook directly from the public CSV snapshot."""
import csv
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1]
COLORS={'UNKNOWN':'E7E6E6','CANDIDATE':'FFF2CC','PARTIAL':'DDEBF7','CONFIRMED':'E2F0D9','OUT_OF_SCOPE':'E7E6E6'}
def read(name):
 with (ROOT/'database'/name).open(newline='') as f:return list(csv.DictReader(f))
def add(wb,name,rows):
 ws=wb.create_sheet(name)
 for row in [list(rows[0])]+[list(r.values()) for r in rows]:
  ws.append(row)
  for cell in ws[ws.max_row]:
   if isinstance(cell.value,str):cell.data_type='s'
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True)
 ws.freeze_panes='B2';ws.sheet_view.zoomScale=85
 for col in ws.columns:
  width=min(65,max(16,max(len(str(c.value or '')) for c in list(col)[:40])+2));ws.column_dimensions[col[0].column_letter].width=width
 for row in ws.iter_rows(min_row=2):
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True)
   if c.value in COLORS:c.fill=PatternFill('solid',fgColor=COLORS[c.value])
  ws.row_dimensions[row[0].row].height=45
 t=Table(displayName='T'+str(len(wb.worksheets)),ref=ws.dimensions);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
 return ws

def main():
 frames=read('can_frames.csv');signals=read('can_signals.csv');obs=read('bridge_observations.csv');wb=Workbook();wb.remove(wb.active)
 counts=Counter(r['Source_ECU'] or 'Unknown / outside scope' for r in frames)
 add(wb,'Overview',[{'Topic':'Vehicle','Result':'Zeekr X, 2023 AWD','Meaning':'Observed vehicle configuration; applicability to other versions requires validation.'},{'Topic':'Observed IDs','Result':str(len(frames)),'Meaning':'Unique frame IDs in the current matrix, not a complete vehicle inventory.'},{'Topic':'Assigned sources','Result':str(sum(bool(r['Source_ECU']) for r in frames)),'Meaning':'Maintainer-confirmed sources; not decoded payloads.'},{'Topic':'Signal rows','Result':str(len(signals)),'Meaning':'4 confirmed VIN layout fields, 35 candidates, 2 unknown fields. No new physical torque calibration.'},{'Topic':'Evidence','Result':'See docs/evidence.md','Meaning':'Raw recordings are not public; hashes identify retained evidence, not independent reproducibility.'}])
 add(wb,'ECU sources',[{'ECU':k,'ID_Count':v,'Share_percent':round(v/len(frames)*100,2),'Meaning':'Transmitter attribution only'} for k,v in counts.items()])
 add(wb,'CAN frames',frames);add(wb,'Signals',signals)
 for name,ecu in [('Rear Drive','RearDrive'),('Front Drive','FrontDrive'),('VCU','VCU'),('BECM','BECM'),('DC-DC','CDD'),('Stalk switch','Stalk switch')]:add(wb,name,[r for r in frames if r['Source_ECU']==ecu])
 add(wb,'Drive role hypotheses',read('drive_role_hypotheses.csv'))
 add(wb,'Bridge evidence',obs)
 add(wb,'Legend',[{'Status':k,'Meaning':v} for k,v in {'UNKNOWN':'Meaning is unknown.','CANDIDATE':'Hypothesis requiring validation.','PARTIAL':'Some fields understood; frame not fully decoded.','CONFIRMED':'Confirmed only within the scope of the column: source or signal.','OUT_OF_SCOPE':'Retained in the inventory; not a current research priority.'}.items()])
 out=ROOT/'reports'/'Zeekr_X_CAN_Matrix_EN.xlsx';out.parent.mkdir(exist_ok=True);wb.save(out);print(out)
if __name__=='__main__':main()
