#!/usr/bin/env python3
"""Validate public data consistency, workbook structure and publication boundaries."""
import csv,re,zipfile
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
def read(name):
 with (ROOT/'database'/name).open(newline='') as f:return list(csv.DictReader(f))
def main():
 frames=read('can_frames.csv');signals=read('can_signals.csv');obs=read('bridge_observations.csv')
 assert len(frames)==len({(r['Bus'],r['CAN_ID']) for r in frames})
 assert Counter(r['Source_ECU'] for r in frames)=={'VCU':46,'BECM':34,'RearDrive':16,'FrontDrive':14,'CDD':11,'Stalk switch':3,'':17}
 assert Counter(r['Status'] for r in signals)=={'CANDIDATE':29,'CONFIRMED':4,'UNKNOWN':2}
 assert Counter(r['Frame_status'] for r in frames)=={'UNKNOWN':121,'CANDIDATE':19,'PARTIAL':1}
 byid={r['CAN_ID']:r for r in frames}
 for r in obs:
  assert byid[r['CAN_ID']]['Source_ECU']==r['Source_ECU']
  if r['Source_ECU']=='RearDrive':assert r['Experiment']=='EXP-0043'
  if r['Source_ECU']=='FrontDrive':assert r['Experiment']=='EXP-0044'
 assert all(not any('channel' in k.lower() for k in r) for r in obs),'Internal capture mapping in public data'
 for r in signals:assert r['Frame_Source_ECU']==byid[r['CAN_ID']]['Source_ECU']
 workbook=ROOT/'reports/Zeekr_X_CAN_Matrix_EN.xlsx'
 with zipfile.ZipFile(workbook) as z:assert z.testzip() is None
 wb=load_workbook(workbook,read_only=True)
 assert wb['CAN frames'].max_row==len(frames)+1
 assert wb['Signals'].max_row==len(signals)+1
 assert wb['Rear Drive'].max_row==17 and wb['Front Drive'].max_row==15
 for ws in wb:
  for row in ws:
   for c in row:
    assert c.data_type!='f','Unexpected spreadsheet formula'
    if isinstance(c.value,str):
     assert not re.search('[\u0400-\u04ff]',c.value),'Non-English workbook cell'
     assert not re.search(r'(?i)\bCAN[12]\b',c.value),'Internal capture mapping in workbook'
 wb.close()
 for folder in ['database','docs','reports']:
  for p in (ROOT/folder).rglob('*'):
   if p.suffix in {'.csv','.md','.json'}:
    text=p.read_text();assert not re.search(r'(?i)\bCAN[12]\b',text),'Internal capture mapping in public text'
    assert not re.search('[\u0400-\u04ff]',text),p
    assert not re.search(r'\b(?:\d{1,3}\.){3}\d{1,3}\b',text),'Unexpected IP address'
    assert not re.search(r'\bL6T[A-HJ-NPR-Z0-9]{14}\b',text),'Unexpected vehicle identifier'
 print('PASS: source groups, statuses, axle assignments, English workbook and publication boundaries')
if __name__=='__main__':main()
